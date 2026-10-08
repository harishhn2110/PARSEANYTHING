"""Document export formats: Markdown, JSON, CSV, and XLSX.

Supports optional PII masking across all output formats.
"""

import csv
from io import BytesIO, StringIO
import json
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill

from app.features.pii import mask_text
from backend.models import Document


class DocumentExporter:
    @classmethod
    def to_markdown(cls, doc: Document, masked: bool = False) -> str:
        fn = doc.filename or "Document"
        conf = int((doc.overall_confidence or 0.95) * 100)
        lines: list[str] = [
            f"# {fn}\n",
            f"> Extracted by ParseAnything Atlas | Confidence: {conf}%\n",
        ]

        # Use pages_rel if populated, otherwise group blocks by page_number
        pages = list(doc.pages_rel) if getattr(doc, "pages_rel", None) else []
        if pages:
            for page in sorted(pages, key=lambda p: (p.page_number if p.page_number is not None else 1)):
                lines.append(f"\n<!-- Page {page.page_number} -->\n")
                page_blocks = list(page.blocks) if getattr(page, "blocks", None) else []
                for b in sorted(page_blocks, key=lambda x: (x.reading_order if x.reading_order is not None else 0.0)):
                    if b.type in ("table_row", "table_cell"):
                        continue

                    raw_text = b.text or ""
                    content = mask_text(raw_text) if masked else raw_text

                    if b.type == "heading":
                        lines.append(f"\n## {content}\n")
                    elif b.type == "list":
                        lines.append(f"- {content}")
                    elif b.type == "table":
                        lines.append(f"\n{content}\n")
                    elif b.type == "equation":
                        lines.append(f"\n$$\n{content}\n$$\n")
                    elif b.type == "figure":
                        lines.append(f"\n![Figure: {b.id}]({content})\n")
                    elif b.type == "caption":
                        lines.append(f"*{content}*\n")
                    elif b.type in ("header", "footer"):
                        lines.append(f"<!-- [{b.type}: {b.id}] {content} -->")
                    else:
                        lines.append(f"\n{content}\n")
        else:
            blocks = list(doc.blocks) if getattr(doc, "blocks", None) else []
            curr_page = None
            for b in sorted(blocks, key=lambda x: (x.page_number or 1, x.reading_order if x.reading_order is not None else 0.0)):
                if b.type in ("table_row", "table_cell"):
                    continue
                if b.page_number != curr_page:
                    curr_page = b.page_number
                    lines.append(f"\n<!-- Page {curr_page} -->\n")

                raw_text = b.text or ""
                content = mask_text(raw_text) if masked else raw_text

                if b.type == "heading":
                    lines.append(f"\n## {content}\n")
                elif b.type == "list":
                    lines.append(f"- {content}")
                elif b.type == "table":
                    lines.append(f"\n{content}\n")
                elif b.type == "equation":
                    lines.append(f"\n$$\n{content}\n$$\n")
                elif b.type == "figure":
                    lines.append(f"\n![Figure: {b.id}]({content})\n")
                elif b.type == "caption":
                    lines.append(f"*{content}*\n")
                elif b.type in ("header", "footer"):
                    lines.append(f"<!-- [{b.type}: {b.id}] {content} -->")
                else:
                    lines.append(f"\n{content}\n")

        return "\n".join(lines).strip()

    @classmethod
    def to_json_dict(cls, doc: Document, masked: bool = False) -> dict[str, Any]:
        pages_data = []
        pages = list(doc.pages_rel) if getattr(doc, "pages_rel", None) else []
        for p in sorted(pages, key=lambda x: (x.page_number if x.page_number is not None else 1)):
            blocks_data = []
            page_blocks = list(p.blocks) if getattr(p, "blocks", None) else []
            for b in sorted(page_blocks, key=lambda x: (x.reading_order if x.reading_order is not None else 0.0)):
                raw_text = b.text or ""
                txt = mask_text(raw_text) if masked else raw_text
                blocks_data.append({
                    "id": b.id,
                    "type": b.type,
                    "page": b.page_number,
                    "reading_order": b.reading_order,
                    "bbox": b.bbox,
                    "text": txt,
                    "confidence": b.confidence,
                    "status": b.status,
                    "extractor": b.extractor,
                    "extractor_version": b.extractor_version,
                    "language": b.language,
                    "flagged": b.flagged,
                    "flag_reason": b.flag_reason,
                    "parent_id": b.parent_id,
                    "data": b.content_json,
                })
            pages_data.append({
                "page": p.page_number,
                "width": p.width,
                "height": p.height,
                "page_type": p.page_type,
                "language": p.language,
                "blocks": blocks_data,
            })

        val_events = [
            {
                "id": ev.id,
                "block_id": ev.block_id,
                "type": ev.type,
                "message": ev.message,
                "severity": ev.severity,
                "expected": ev.expected,
                "actual": ev.actual,
                "difference": ev.difference,
                "status": ev.status,
                "created_at": ev.created_at.isoformat() if ev.created_at else None,
            }
            for ev in (getattr(doc, "validation_events", None) or [])
        ]

        return {
            "document_id": doc.id,
            "filename": doc.filename,
            "format": doc.format,
            "size_bytes": doc.size_bytes,
            "pages_count": doc.pages,
            "status": doc.status,
            "overall_confidence": doc.overall_confidence,
            "metrics": doc.metrics_json or {},
            "pages": pages_data,
            "validation_events": val_events,
        }

    @classmethod
    def to_csv(cls, doc: Document, masked: bool = False) -> str:
        out = StringIO()
        writer = csv.writer(out)

        all_blocks = getattr(doc, "blocks", None) or []
        table_blocks = [b for b in all_blocks if getattr(b, "type", None) == "table"]
        if not table_blocks:
            writer.writerow(["No tables extracted from this document."])
            return out.getvalue()

        for t_idx, tb in enumerate(table_blocks, 1):
            writer.writerow([f"=== Table {t_idx} (Block ID: {tb.id}, Page: {tb.page_number}) ==="])
            d = tb.content_json or {}
            headers = d.get("headers", [])
            rows = d.get("rows", [])
            if headers:
                writer.writerow([mask_text(str(h)) if masked else str(h) for h in headers])
            for r in rows:
                writer.writerow([mask_text(str(c)) if masked else str(c) for c in r])
            writer.writerow([])

        return out.getvalue()

    @classmethod
    def to_xlsx(cls, doc: Document, masked: bool = False) -> bytes:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        all_blocks = getattr(doc, "blocks", None) or []
        table_blocks = [b for b in all_blocks if getattr(b, "type", None) == "table"]
        if not table_blocks:
            ws = wb.create_sheet(title="Tables")
            ws.cell(row=1, column=1, value="No tables extracted from this document.")
        else:
            for t_idx, tb in enumerate(table_blocks, 1):
                sheet_title = f"Table_{t_idx}_{tb.id}"[:31]
                ws = wb.create_sheet(title=sheet_title)

                d = tb.content_json or {}
                headers = d.get("headers", [])
                rows = d.get("rows", [])

                cur_row = 1
                ws.cell(row=cur_row, column=1, value=f"Table: {tb.id} (Page {tb.page_number})").font = Font(bold=True)
                cur_row += 1

                if headers:
                    for c_idx, h in enumerate(headers, 1):
                        cell = ws.cell(row=cur_row, column=c_idx, value=mask_text(str(h)) if masked else str(h))
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
                    cur_row += 1

                for row in rows:
                    for c_idx, val in enumerate(row, 1):
                        ws.cell(row=cur_row, column=c_idx, value=mask_text(str(val)) if masked else str(val))
                    cur_row += 1

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
